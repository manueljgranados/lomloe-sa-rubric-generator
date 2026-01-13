from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path("src/lomloe_sa_gen/assets/templates/sa_template.docx")


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # Portada
    p = doc.add_paragraph("{{ titulo }}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = "Title"

    p = doc.add_paragraph("{{ materia }} · {{ nivel }}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")  # espacio
    p = doc.add_paragraph("Resumen:")
    p.runs[0].bold = True
    doc.add_paragraph("{{ resumen }}")

    p = doc.add_paragraph("Generado el {{ fecha }}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Pie de página (si existe)
    section = doc.sections[0]
    footer = section.footer
    footer.paragraphs[0].text = "{{ footer_text }}"

    doc.add_page_break()

    # Cuerpo
    doc.add_heading("Situación de Aprendizaje", level=1)

    doc.add_heading("Competencias específicas", level=2)
    doc.add_paragraph("{% for c in competencias %}", style="Normal")
    doc.add_paragraph("• {{ c }}", style="Normal")
    doc.add_paragraph("{% endfor %}", style="Normal")

    doc.add_heading("Criterios de evaluación", level=2)
    doc.add_paragraph("{% for c in criterios %}", style="Normal")
    doc.add_paragraph("• {{ c }}", style="Normal")
    doc.add_paragraph("{% endfor %}", style="Normal")

    doc.add_heading("Producto final", level=2)
    doc.add_paragraph("{{ producto_final }}")

    doc.add_heading("Metodología", level=2)
    doc.add_paragraph("{{ metodologia }}")

    doc.add_heading("Atención a la diversidad", level=2)
    doc.add_paragraph("{{ atencion_diversidad }}")

    doc.add_heading("Instrumentos de evaluación", level=2)
    doc.add_paragraph("{% for i in instrumentos_evaluacion %}", style="Normal")
    doc.add_paragraph("• {{ i }}", style="Normal")
    doc.add_paragraph("{% endfor %}", style="Normal")

    doc.add_heading("Secuenciación de sesiones", level=2)
    doc.add_paragraph("{% for s in sesiones %}", style="Normal")
    doc.add_heading("Sesión {{ loop.index }}. {{ s.title }} ({{ s.minutes }} min)", level=3)
    doc.add_paragraph("Actividades: {{ s.activities }}")
    doc.add_paragraph("Evidencia: {{ s.evidence }}")
    doc.add_paragraph("{% endfor %}", style="Normal")

    doc.add_page_break()
    doc.add_heading("Rúbrica (borrador)", level=1)

    # Tabla de rúbrica (docxtpl: bucle en filas)
    table = doc.add_table(rows=2, cols=5)
    hdr = table.rows[0].cells
    hdr[0].text = "Criterio"
    hdr[1].text = "{{ rubric_levels[0] }}"
    hdr[2].text = "{{ rubric_levels[1] }}"
    hdr[3].text = "{{ rubric_levels[2] }}"
    hdr[4].text = "{{ rubric_levels[3] }}"

    row = table.rows[1].cells
    row[0].text = "{% for row in rubric_rows %}{{ row.criterion }}"
    row[1].text = "{{ row.descriptors[0] }}"
    row[2].text = "{{ row.descriptors[1] }}"
    row[3].text = "{{ row.descriptors[2] }}"
    row[4].text = "{{ row.descriptors[3] }}{% endfor %}"

    doc.save(OUT)
    print(f"Created: {OUT}")


if __name__ == "__main__":
    main()
