from __future__ import annotations

import io

from docx import Document

from lomloe_sa_gen.core.models import SASpec


def export_sa_docx(spec: SASpec) -> bytes:
    doc = Document()

    doc.add_heading(spec.titulo, level=1)

    p = doc.add_paragraph()
    p.add_run("Nivel: ").bold = True
    p.add_run(spec.nivel)

    p = doc.add_paragraph()
    p.add_run("Materia: ").bold = True
    p.add_run(spec.materia)

    doc.add_heading("Competencias específicas", level=2)
    for c in spec.competencias:
        doc.add_paragraph(c, style="List Bullet")

    doc.add_heading("Criterios de evaluación", level=2)
    for c in spec.criterios:
        doc.add_paragraph(c, style="List Bullet")

    doc.add_heading("Producto final", level=2)
    doc.add_paragraph(spec.producto_final)

    doc.add_heading("Metodología", level=2)
    doc.add_paragraph(spec.metodologia)

    doc.add_heading("Atención a la diversidad", level=2)
    doc.add_paragraph(spec.atencion_diversidad)

    doc.add_heading("Instrumentos de evaluación", level=2)
    for i in spec.instrumentos_evaluacion:
        doc.add_paragraph(i, style="List Bullet")

    doc.add_heading("Secuenciación de sesiones", level=2)
    for idx, s in enumerate(spec.sesiones, start=1):
        doc.add_heading(f"Sesión {idx}. {s.title} ({s.minutes} min)", level=3)
        doc.add_paragraph(f"Actividades: {s.activities}")
        doc.add_paragraph(f"Evidencia: {s.evidence}")

    # Rúbrica (tabla mínima)
    doc.add_page_break()
    doc.add_heading("Rúbrica (borrador)", level=1)

    levels = ["4 - Excelente", "3 - Notable", "2 - Suficiente", "1 - En proceso"]
    table = doc.add_table(rows=1, cols=1 + len(levels))
    hdr = table.rows[0].cells
    hdr[0].text = "Criterio"
    for j, lvl in enumerate(levels, start=1):
        hdr[j].text = lvl

    for c in spec.criterios:
        row = table.add_row().cells
        row[0].text = c
        for j in range(1, 1 + len(levels)):
            row[j].text = "Definir descriptor."

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
